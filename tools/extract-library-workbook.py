import json
import re
import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook


def clean(value):
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value).strip()


def safe_code(value):
    return re.sub(r"[^a-z0-9_-]+", "-", clean(value).lower()).strip("-")


def image_position(sheet, image):
    anchor = image.anchor._from
    default_height = sheet.sheet_format.defaultRowHeight or 15
    rows_height = sum(
        (sheet.row_dimensions[row].height or default_height) * 12700
        for row in range(1, anchor.row + 1)
    )
    return rows_height + anchor.rowOff, anchor.col, anchor.colOff


source = Path(sys.argv[1])
output_root = Path(sys.argv[2])
covers_root = output_root / "assets" / "portadas"
data_root = output_root / "data"
covers_root.mkdir(parents=True, exist_ok=True)
data_root.mkdir(parents=True, exist_ok=True)

workbook = load_workbook(source, read_only=False, data_only=True)
sheet = workbook.active
book_rows = [row for row in range(11, sheet.max_row + 1) if clean(sheet.cell(row, 1).value)]
images = sorted(sheet._images, key=lambda image: image_position(sheet, image))

if len(book_rows) != len(images):
    raise RuntimeError(f"Se encontraron {len(book_rows)} libros y {len(images)} portadas; no es seguro relacionarlos.")

books = []
for row, image in zip(book_rows, images):
    code = clean(sheet.cell(row, 1).value)
    document_id = safe_code(code)
    image_format = (image.format or "png").lower()
    extension = "jpg" if image_format in {"jpg", "jpeg"} else image_format
    cover_name = f"{document_id}.{extension}"
    (covers_root / cover_name).write_bytes(image._data())

    source_status = clean(sheet.cell(row, 12).value).upper()
    explicitly_unavailable = source_status in {"OCUPADO", "NO ESTA EN BIBLIOTECA"}
    books.append({
        "id": document_id,
        "sourceCode": code,
        "title": clean(sheet.cell(row, 2).value),
        "author": clean(sheet.cell(row, 5).value),
        "category": clean(sheet.cell(row, 6).value) or "General",
        "review": clean(sheet.cell(row, 4).value),
        "isbn": "",
        "editionYear": clean(sheet.cell(row, 7).value),
        "publisher": clean(sheet.cell(row, 8).value),
        "country": clean(sheet.cell(row, 9).value),
        "location": clean(sheet.cell(row, 10).value),
        "pages": int(sheet.cell(row, 11).value) if isinstance(sheet.cell(row, 11).value, (int, float)) else None,
        "sourceStatus": source_status or "DISPONIBLE",
        "availableOn": clean(sheet.cell(row, 13).value),
        "coverUrl": f"./assets/portadas/{cover_name}",
        "totalCopies": 1,
        "availableCopies": 0 if explicitly_unavailable else 1,
        "active": True,
    })

output_file = data_root / "books-import.json"
output_file.write_text(json.dumps({"source": source.name, "count": len(books), "books": books}, ensure_ascii=False, indent=2), encoding="utf-8")
print(json.dumps({"books": len(books), "covers": len(images), "output": str(output_file)}, ensure_ascii=False))
